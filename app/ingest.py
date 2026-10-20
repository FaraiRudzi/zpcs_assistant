"""Objective 1: extract and clean text, then split it into citable chunks."""
import re
from pathlib import Path
from .config import CHUNK_WORDS, CHUNK_OVERLAP, OCR_ENABLED

def clean_text(t: str) -> str:
    t = t.replace("\r", "\n")
    t = re.sub(r"-\n(?=[a-z])", "", t)      # join words hyphenated across lines
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()

def is_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 160:
        return False
    if re.match(r"^(section|article|part|chapter|schedule)\s+[\w.\-]+", line, re.I):
        return True
    return bool(re.match(r"^\d+[A-Z]{0,2}(\.\d+)*[.)]?\s+[A-Z]", line)) and not line.endswith((".", ",", ";"))

def _ocr_pdf_page(path: Path, page_no: int) -> str:
    try:
        from pdf2image import convert_from_path
        import pytesseract
        img = convert_from_path(str(path), first_page=page_no, last_page=page_no, dpi=200)[0]
        return pytesseract.image_to_string(img)
    except Exception:
        return ""

def extract_pages(path: Path):
    """Return a list of (page_number, text). Page 0 means page numbers are unavailable."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        pages = []
        for i, p in enumerate(PdfReader(str(path)).pages, 1):
            txt = p.extract_text() or ""
            if len(txt.strip()) < 30 and OCR_ENABLED:
                txt = _ocr_pdf_page(path, i)
            pages.append((i, clean_text(txt)))
        return pages
    if ext == ".doc":
        import subprocess, tempfile
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(["soffice", "--headless", "--convert-to", "docx", "--outdir", tmp, str(path)],
                           check=True, capture_output=True, timeout=180)
            return extract_pages(Path(tmp) / (path.stem + ".docx"))
    if ext == ".docx":
        from docx import Document
        d = Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return [(0, clean_text("\n".join(parts)))]
    if ext == ".txt":
        return [(0, clean_text(path.read_text(encoding="utf-8", errors="ignore")))]
    raise ValueError("Unsupported file type")

def is_toc_page(text: str) -> bool:
    """A contents page is mostly short numbered titles with no sentences. Skipping it stops the
    system from 'finding' a section title instead of the section itself."""
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    if len(lines) < 8:
        return False
    titles = sum(1 for l in lines
                 if len(l.split()) <= 16 and not l.endswith((",", ";"))
                 and re.match(r"^(\d+[A-Z]?\.?|PART\b|[A-Z]\.|\([a-z]\))\s*\S", l))
    dotted = sum(1 for l in lines if re.search(r"\.{4,}\s*\d+$", l))
    return (titles + dotted) / len(lines) > 0.6

def chunk_pages(pages):
    """One chunk never mixes two sections. Every chunk starts with its section heading so that
    both search and citations point to the right section (e.g. '39 Assistance by private persons...')."""
    chunks, heading = [], ""
    step = max(1, CHUNK_WORDS - CHUNK_OVERLAP)
    for page, text in pages:
        if is_toc_page(text):
            continue
        segments, cur, seg_heading = [], [], heading
        for line in text.split("\n"):
            if is_heading(line):
                if cur:
                    segments.append((seg_heading, cur))
                    cur = []
                heading = seg_heading = re.sub(r"\s+", " ", line.strip())
            else:
                cur.extend(line.split())
        if cur:
            segments.append((seg_heading, cur))
        for h, words in segments:
            i = 0
            while i < len(words):
                body = " ".join(words[i:i + CHUNK_WORDS])
                chunks.append({"page": page, "section": h, "text": (h + ". " if h else "") + body})
                if i + CHUNK_WORDS >= len(words):
                    break
                i += step
    return chunks
